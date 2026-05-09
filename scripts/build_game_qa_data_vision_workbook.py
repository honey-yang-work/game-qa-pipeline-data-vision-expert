from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from validate_game_qa_payload import validate_payload


CASE_COLUMNS = [
    "模块",
    "ID",
    "子功能",
    "标题",
    "优先级(P0-P3)",
    "追踪编号",
    "前置条件",
    "步骤",
    "预期结果",
]

RTM_COLUMNS = [
    "模块",
    "追踪编号",
    "追踪类型",
    "追踪描述",
    "测试用例ID",
    "覆盖状态",
    "备注",
]

CODE_TRACE_COLUMNS = [
    "需求ID",
    "表字段ID",
    "配置表路径",
    "Sheet",
    "配置字段",
    "配置值",
    "默认值",
    "边界值",
    "配置风险",
    "协议/入口",
    "业务承载",
    "数据承载",
    "输出承载",
    "风险点",
    "测试用例ID",
    "命中置信度",
]

CASE_ALIASES = {
    "模块": ["模块"],
    "ID": ["ID"],
    "子功能": ["子功能"],
    "标题": ["标题"],
    "优先级(P0-P3)": ["优先级(P0-P3)"],
    "追踪编号": ["追踪编号", "需求编号", "关联编号"],
    "前置条件": ["前置条件"],
    "步骤": ["步骤"],
    "预期结果": ["预期结果"],
}

RTM_ALIASES = {
    "模块": ["模块"],
    "追踪编号": ["追踪编号", "需求编号", "关联编号"],
    "追踪类型": ["追踪类型", "编号类型"],
    "追踪描述": ["追踪描述", "需求描述", "字段描述", "需求/字段描述"],
    "测试用例ID": ["测试用例ID"],
    "覆盖状态": ["覆盖状态"],
    "备注": ["备注"],
}

CODE_TRACE_ALIASES = {
    "需求ID": ["需求ID", "需求编号", "R"],
    "表字段ID": ["表字段ID", "字段ID", "T"],
    "配置表路径": ["配置表路径", "表格路径", "配置文件", "配置表"],
    "Sheet": ["Sheet", "工作表", "页签"],
    "配置字段": ["配置字段", "字段名", "配置列"],
    "配置值": ["配置值", "字段值", "命中值"],
    "默认值": ["默认值", "Default"],
    "边界值": ["边界值", "最小值/最大值", "范围"],
    "配置风险": ["配置风险", "配置问题", "配置缺陷"],
    "协议/入口": ["协议/入口", "入口", "协议", "路由", "RPC"],
    "业务承载": ["业务承载", "业务代码", "Service", "Handler", "Controller"],
    "数据承载": ["数据承载", "配置/DB/Cache", "配置", "DB", "Cache"],
    "输出承载": ["输出承载", "回包/事件/日志", "回包", "事件", "日志"],
    "风险点": ["风险点", "风险"],
    "测试用例ID": ["测试用例ID"],
    "命中置信度": ["命中置信度", "置信度"],
}

HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")
HEADER_FONT = Font(bold=True)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build a game QA workbook for text, visual, and table-driven traceability."
    )
    parser.add_argument("--input-json", required=True, help="Path to the input JSON file.")
    parser.add_argument("--out", required=True, help="Path to the output .xlsx file.")
    parser.add_argument("--system-name", default="", help="Optional expected system-name prefix.")
    parser.add_argument("--skip-validate", action="store_true", help="Skip payload validation before export.")
    return parser.parse_args()


def load_payload(path: Path) -> dict:
    with path.open("r", encoding="utf-8-sig") as handle:
        payload = json.load(handle)

    if not isinstance(payload, dict):
        raise ValueError("Input JSON must be an object with test_cases and rtm arrays.")

    return payload


def ensure_list(value, field_name: str) -> list:
    if value is None:
        return []
    if not isinstance(value, list):
        raise ValueError(f"Input JSON field {field_name} must be a list.")
    return value


def normalize_scalar(value) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return ", ".join(str(item) for item in value)
    return str(value)


def lookup_with_aliases(record: dict, aliases: dict[str, list[str]], canonical_key: str, fallback: str = "") -> str:
    for alias in aliases[canonical_key]:
        if alias in record:
            return normalize_scalar(record[alias])
    return fallback


def normalize_records(
    records: Iterable[dict],
    columns: list[str],
    aliases: dict[str, list[str]],
    module_name: str = "",
) -> list[dict]:
    normalized = []
    for record in records:
        if not isinstance(record, dict):
            raise ValueError("Every row in test_cases or rtm must be an object.")

        normalized_row = {}
        for column in columns:
            fallback = module_name if column == "模块" else ""
            normalized_row[column] = lookup_with_aliases(record, aliases, column, fallback)
        normalized.append(normalized_row)
    return normalized


def write_sheet(writer: pd.ExcelWriter, sheet_name: str, rows: list[dict], columns: list[str]) -> None:
    frame = pd.DataFrame(rows, columns=columns)
    frame.to_excel(writer, index=False, sheet_name=sheet_name)


def autosize_and_style(writer: pd.ExcelWriter) -> None:
    for worksheet in writer.book.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL

        for column_index, column_cells in enumerate(worksheet.columns, start=1):
            max_length = 0
            for cell in column_cells:
                text = "" if cell.value is None else str(cell.value)
                for line in text.splitlines() or [""]:
                    max_length = max(max_length, len(line))
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 12), 60)


def build_dataset(payload: dict) -> tuple[list[dict], list[dict], list[dict]]:
    test_cases = normalize_records(
        ensure_list(payload.get("test_cases"), "test_cases"),
        CASE_COLUMNS,
        CASE_ALIASES,
    )
    rtm_rows = normalize_records(
        ensure_list(payload.get("rtm"), "rtm"),
        RTM_COLUMNS,
        RTM_ALIASES,
    )
    code_trace_rows = normalize_records(
        ensure_list(payload.get("code_traceability"), "code_traceability"),
        CODE_TRACE_COLUMNS,
        CODE_TRACE_ALIASES,
    )
    return test_cases, rtm_rows, code_trace_rows


def build_workbook(payload: dict, output_path: Path, system_name: str = "", skip_validate: bool = False) -> Path:
    if not skip_validate:
        errors = validate_payload(payload, system_name)
        if errors:
            formatted_errors = "\n".join(f"- {error}" for error in errors)
            raise ValueError(f"Payload validation failed:\n{formatted_errors}")

    test_cases, rtm_rows, code_trace_rows = build_dataset(payload)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        write_sheet(writer, "测试用例库", test_cases, CASE_COLUMNS)
        write_sheet(writer, "需求跟踪矩阵", rtm_rows, RTM_COLUMNS)
        if code_trace_rows:
            write_sheet(writer, "代码追踪矩阵", code_trace_rows, CODE_TRACE_COLUMNS)
        autosize_and_style(writer)

    return output_path.resolve()


def main() -> int:
    args = parse_args()
    payload = load_payload(Path(args.input_json))
    output_path = build_workbook(payload, Path(args.out), args.system_name, args.skip_validate)
    print(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
